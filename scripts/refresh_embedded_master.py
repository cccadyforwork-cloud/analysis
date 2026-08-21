import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook


def main() -> None:
    html_path = Path(sys.argv[1])
    workbook_path = Path(sys.argv[2])
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(rows)]
    master_rows = []
    for values in rows:
        record = {headers[i]: values[i] for i in range(min(len(headers), len(values))) if headers[i]}
        if str(record.get("ASIN") or "").strip():
            master_rows.append(record)

    html = html_path.read_text(encoding="utf-8")
    pattern = re.compile(r'(<script id="wk34-embedded-data" type="application/json">)(.*?)(</script>)', re.S)
    match = pattern.search(html)
    if not match:
        raise RuntimeError("embedded data block not found")
    data = json.loads(match.group(2))
    data["masterRows"] = master_rows
    payload = json.dumps(data, ensure_ascii=True, separators=(",", ":"), default=str)
    html = html[: match.start()] + match.group(1) + payload + match.group(3) + html[match.end() :]
    html_path.write_text(html, encoding="utf-8", newline="\n")
    print(json.dumps({"rows": len(master_rows), "columns": headers}, ensure_ascii=False))


if __name__ == "__main__":
    main()
