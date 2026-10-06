"""Embed ASIN → 品名 mappings from a Listing workbook and optional SKU catalog."""

import argparse
import json
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_HTML = ROOT / "货盘动销分析_WK29-WK30_可下载版.html"


def read_sku_catalog(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    names_by_sku = defaultdict(set)
    names_by_model = defaultdict(set)
    for sheet in workbook:
        rows = iter(sheet.values)
        headers = [str(value or "").strip() for value in next(rows)]
        if "*SKU" not in headers or "品名" not in headers:
            continue
        sku_col, name_col = headers.index("*SKU"), headers.index("品名")
        model_col = headers.index("型号") if sheet.title == "产品" and "型号" in headers else -1
        for row in rows:
            sku = str(row[sku_col] or "").strip().upper()
            name = str(row[name_col] or "").strip()
            if not sku or not name:
                continue
            names_by_sku[sku].add(name)
            if model_col >= 0:
                model = str(row[model_col] or "").strip().upper()
                if model:
                    names_by_model[model].add(name)
    conflicts = [sku for sku, names in names_by_sku.items() if len(names) > 1]
    if conflicts:
        raise ValueError(f"Conflicting product names for SKU: {', '.join(conflicts[:5])}")
    exact = {sku: next(iter(names)) for sku, names in names_by_sku.items()}
    models = {model: next(iter(names)) for model, names in names_by_model.items() if len(names) == 1}
    return exact, models


def sku_name(sku, exact, models):
    sku = sku.strip().upper()
    if sku in exact:
        return exact[sku]
    # QS size variants share a model but some sizes are absent from the catalog.
    # Inherit a model name only when its SKU suffix is a numeric size.
    if sku.startswith("QS-"):
        matches = [
            (model, name)
            for model, name in models.items()
            if model.startswith("QS-")
            and (sku == model or (
                sku.startswith(model + "-")
                and re.fullmatch(r"\d{2}(?:/|-)\d{2}", sku[len(model) + 1:])
            ))
        ]
        if matches:
            longest = max(len(model) for model, _ in matches)
            names = {name for model, name in matches if len(model) == longest}
            if len(names) == 1:
                return next(iter(names))
    return ""


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path, help="Listing .xlsx with ASIN and 品名 columns")
    parser.add_argument("--sku-workbook", type=Path, help="Product catalog .xlsx with *SKU, 品名 and 型号")
    parser.add_argument("--html", type=Path, default=DEFAULT_HTML)
    parser.add_argument("--date", help="Data date YYYY-MM-DD; defaults to Listing filename date")
    args = parser.parse_args()

    date_match = re.search(r"(20\d{2})(\d{2})(\d{2})", args.workbook.name)
    date = args.date or ("-".join(date_match.groups()) if date_match else "")
    exact, models = read_sku_catalog(args.sku_workbook) if args.sku_workbook else ({}, {})

    workbook = load_workbook(args.workbook, read_only=True, data_only=True)
    rows = iter(workbook.active.values)
    headers = [str(value or "").strip() for value in next(rows)]
    asin_col, name_col = headers.index("ASIN"), headers.index("品名")
    sku_col = headers.index("MSKU") if "MSKU" in headers else -1
    parent_col = headers.index("父ASIN") if "父ASIN" in headers else -1
    names = {}
    sku_candidates = defaultdict(set)
    exact_candidates = defaultdict(set)
    children = defaultdict(set)
    for row in rows:
        asin = str(row[asin_col] or "").strip().upper()
        name = str(row[name_col] or "").strip()
        if not re.fullmatch(r"[A-Z0-9]{10}", asin):
            continue
        if name:
            if asin in names and names[asin] != name:
                raise ValueError(f"Conflicting names for {asin}: {names[asin]!r} / {name!r}")
            names[asin] = name
        if sku_col >= 0 and exact:
            msku = str(row[sku_col] or "").strip().upper()
            candidate = sku_name(msku, exact, models)
            if candidate:
                sku_candidates[asin].add(candidate)
            if msku in exact:
                exact_candidates[asin].add(exact[msku])
        if parent_col >= 0:
            parent = str(row[parent_col] or "").strip().upper()
            if re.fullmatch(r"[A-Z0-9]{10}", parent) and parent != asin:
                children[parent].add(asin)

    direct_count = len(names)
    corrected = []
    for asin, candidates in exact_candidates.items():
        if asin in names and len(candidates) == 1:
            catalog_name = next(iter(candidates))
            if names[asin] != catalog_name:
                corrected.append((asin, names[asin], catalog_name))
                names[asin] = catalog_name
    for asin, candidates in sku_candidates.items():
        if asin not in names and len(candidates) == 1:
            names[asin] = next(iter(candidates))
    sku_count = len(names) - direct_count
    # Parent names are inherited only when every listed child has the same name.
    for parent, child_asins in children.items():
        if parent in names:
            continue
        child_names = [names.get(asin) for asin in child_asins]
        if child_names and all(child_names) and len(set(child_names)) == 1:
            names[parent] = child_names[0]
    parent_count = len(names) - direct_count - sku_count

    html = args.html.read_text(encoding="utf-8")
    payload = json.dumps(names, ensure_ascii=True, separators=(",", ":")).replace("<", "\\u003c")
    pattern = r"var PRODUCT_NAMES = .*?;\nvar PRODUCT_NAMES_DATE = .*?;"
    replacement = f"var PRODUCT_NAMES = {payload};\nvar PRODUCT_NAMES_DATE = {json.dumps(date)};"
    updated, count = re.subn(pattern, lambda _: replacement, html, count=1)
    if count != 1:
        raise ValueError("PRODUCT_NAMES marker not found exactly once in HTML")
    args.html.write_text(updated, encoding="utf-8")
    print(f"Embedded {len(names)} names: {direct_count} direct, {sku_count} by SKU/model, {parent_count} by parent ASIN.")
    for asin, old_name, new_name in corrected:
        print(f"Corrected {asin} from {old_name!r} to {new_name!r} using exact MSKU catalog match.")


if __name__ == "__main__":
    main()
