import json
import sys
import time
from pathlib import Path

from openpyxl import load_workbook


HEADER = ["父ASIN", "子ASIN", "标题", "会话数总计", "", "转化率总计", "", "页面浏览量总计", "", "", "", "", "", "已订购商品数量", "", "", "", "", "已订购商品销售额"]


def quote(value):
    return '"%s"' % str(value if value is not None else "").replace('"', '""')


def normalize(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        source = workbook.active.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(source)]
        indexes = {header: index for index, header in enumerate(headers)}
        lines = [",".join(quote(value) for value in HEADER)]
        for values in source:
            asin = str(values[indexes["ASIN"]] or "").strip().upper()
            if not asin:
                continue
            name = values[indexes["品名"]] or ""
            units = values[indexes["数量"]] or 0
            sales = values[indexes["销售额(Item Price)"]] or 0
            row = [asin, asin, name, 0, "", 0, "", 0, "", "", "", "", "", units, "", "", "", "", sales]
            lines.append(",".join(quote(value) for value in row))
    finally:
        workbook.close()
    return {"store": "全部店铺", "name": path.name, "text": "\n".join(lines)}


def main():
    output = Path(sys.argv[1])
    history_paths = [Path(value) for value in sys.argv[2:-1]]
    current_path = Path(sys.argv[-1])
    now = int(time.time() * 1000)
    state = {
        "history": {"record": [normalize(path) for path in history_paths], "file_name": "、".join(path.name for path in history_paths), "uploaded_at": now},
        "current": {"record": normalize(current_path), "file_name": current_path.name, "uploaded_at": now},
    }
    output.write_text(json.dumps(state, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(json.dumps({"history": len(history_paths), "current": current_path.name, "bytes": output.stat().st_size}, ensure_ascii=False))


if __name__ == "__main__":
    main()
